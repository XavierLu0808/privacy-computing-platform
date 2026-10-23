#!/usr/bin/env python3
"""Markdown 转 Word（宋体正式文档版式）。

支持：# 至 #### 标题、段落、有序/无序列表、GFM 表格、围栏代码块、图片 ![说明](路径)、
行内 **加粗** 与 `代码`、文首 YAML 风格 front matter（封面信息）。

版式：A4，全文宋体（西文 Times New Roman），正文小四 12 pt、1.5 倍行距、首行缩进 2 字；
标题加粗黑色；表格带边框、表头浅灰底并跨页重复；封面 + 目录页（Word TOC 域）+ 页码。

用法：
  python3 tools/md2docx.py 输入.md 输出.docx
  python3 tools/md2docx.py 输入.md 输出.docx --title 项目名 --subtitle 文档名 --version V1.0 --date "2026 年 10 月" --org 复旦大学

封面字段优先级：命令行参数 > md 文首 front matter（title / subtitle / version / date / org）。
"""
import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = '宋体'
LATIN_FONT = 'Times New Roman'
MONO_FONT = 'Courier New'
BODY_PT = 12          # 小四
TABLE_PT = 10.5       # 五号
CODE_PT = 9
TEXT_WIDTH_CM = 15.4  # A4 宽 21 cm，左右页边距各 2.8 cm
HEAD_SIZES = {1: 16, 2: 14, 3: 12, 4: 12}   # 三号、四号、小四、小四


# ---------------------------------------------------------------- 字体与样式

def set_fonts(rpr, ascii_font=LATIN_FONT, east=CN_FONT):
    """在 rPr 上写 rFonts，并去掉主题字体属性（否则 Word 优先用主题字体）。"""
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    rfonts.set(qn('w:ascii'), ascii_font)
    rfonts.set(qn('w:hAnsi'), ascii_font)
    rfonts.set(qn('w:eastAsia'), east)
    rfonts.set(qn('w:cs'), ascii_font)


def style_rpr(style):
    return style.element.get_or_add_rPr()


def setup_styles(doc):
    styles = doc.styles
    # 文档默认
    rpr_default = styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    set_fonts(rpr_default)
    normal = styles['Normal']
    set_fonts(style_rpr(normal))
    normal.font.size = Pt(BODY_PT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for level, size in HEAD_SIZES.items():
        st = styles['Heading %d' % level]
        set_fonts(style_rpr(st))
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        hpf = st.paragraph_format
        hpf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        hpf.space_before = Pt({1: 18, 2: 12, 3: 8, 4: 6}[level])
        hpf.space_after = Pt({1: 10, 2: 6, 3: 4, 4: 2}[level])
        hpf.keep_with_next = True
        hpf.first_line_indent = Pt(0)
    for name in ('Header', 'Footer'):
        set_fonts(style_rpr(styles[name]))
        styles[name].font.size = Pt(9)


def run_fmt(run, size=None, bold=None, mono=False):
    rpr = run._element.get_or_add_rPr()
    if mono:
        set_fonts(rpr, ascii_font=MONO_FONT)
    else:
        set_fonts(rpr)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    return run


# ---------------------------------------------------------------- 行内解析

INLINE_RX = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def unescape(text):
    return re.sub(r'\\([|*`\\])', r'\1', text)


def add_inline(par, text, size=None, bold=None):
    for part in INLINE_RX.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run_fmt(par.add_run(unescape(part[2:-2])), size=size, bold=True)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            run_fmt(par.add_run(unescape(part[1:-1])), size=(size or BODY_PT) - 1, bold=bold, mono=True)
        else:
            run_fmt(par.add_run(unescape(part)), size=size, bold=bold)


# ---------------------------------------------------------------- 表格

def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def set_table_borders(table):
    tblpr = table._element.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblpr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge, w in (('top', 40), ('left', 90), ('bottom', 40), ('right', 90)):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:w'), str(w))
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    tblpr.append(margins)


def disp_len(s):
    """显示宽度：中日韩字符与全角符号按 2 计，其余按 1 计。"""
    return sum(2 if ord(c) > 0x2E80 else 1 for c in s)


def col_widths(rows, ncols):
    weights = []
    for j in range(ncols):
        lens = [disp_len(re.sub(r'[*`\\]', '', r[j])) for r in rows if j < len(r)]
        mx = max(lens) if lens else 4
        avg = sum(lens) / len(lens) if lens else 4
        weights.append(max(6.0, min(mx, 70) * 0.45 + avg * 0.35 + 3))
    total = sum(weights)
    widths = [TEXT_WIDTH_CM * w / total for w in weights]
    # 每列下限：最长不可折行片段（英文标识符等）的宽度，五号字半角约 0.19 cm
    mins = []
    for j in range(ncols):
        tok = max((disp_len(t) for r in rows if j < len(r)
                   for t in re.split(r'[\s，。；、（）/]+', re.sub(r'[*`\\]', '', r[j]))), default=0)
        mins.append(max(1.9, min(tok * 0.24 + 0.5, 5.0)))
    # 列多时下限之和可能超过版心宽度：等比缩小下限（每列不小于 1.2 cm）
    if sum(mins) > TEXT_WIDTH_CM * 0.95:
        k = TEXT_WIDTH_CM * 0.95 / sum(mins)
        mins = [max(1.2, m * k) for m in mins]
    # 过窄的列顶到下限，差额从高于下限的列按富余量比例扣除
    for _ in range(ncols):
        short = [i for i, w in enumerate(widths) if w < mins[i] - 1e-6]
        if not short:
            break
        need = sum(mins[i] - widths[i] for i in short)
        for i in short:
            widths[i] = mins[i]
        spare = {i: widths[i] - mins[i] for i in range(ncols) if i not in short and widths[i] > mins[i]}
        tot = sum(spare.values())
        if tot <= 0:
            break
        for i, s in spare.items():
            widths[i] -= need * s / tot
    return widths


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    rows = [r + [''] * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    tblpr = table._element.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblpr.append(layout)
    widths = col_widths(rows, ncols)
    for j, w in enumerate(widths):
        table.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        tr = table.rows[i]
        trpr = tr._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        trpr.append(cant)
        if i == 0:
            hdr = OxmlElement('w:tblHeader')
            trpr.append(hdr)
        for j, text in enumerate(row):
            cell = tr.cells[j]
            cell.width = Cm(widths[j])
            par = cell.paragraphs[0]
            pf = par.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)
            if i == 0:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(par, text, size=TABLE_PT, bold=True)
                shade(cell, 'D9D9D9')
            else:
                add_inline(par, text, size=TABLE_PT)
    # 表后留一点间距
    gap = doc.add_paragraph()
    gap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    gap.paragraph_format.space_after = Pt(4)
    for r in gap.runs:
        r.font.size = Pt(4)
    return table


def split_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|') and not line.endswith('\\|'):
        line = line[:-1]
    cells, cur, i = [], '', 0
    while i < len(line):
        c = line[i]
        if c == '\\' and i + 1 < len(line) and line[i + 1] == '|':
            cur += '\\|'
            i += 2
            continue
        if c == '|':
            cells.append(cur.strip())
            cur = ''
        else:
            cur += c
        i += 1
    cells.append(cur.strip())
    return cells


def is_sep_row(line):
    cells = split_row(line)
    return bool(cells) and all(re.fullmatch(r':?-{3,}:?', c) for c in cells)


# ---------------------------------------------------------------- 段落类元素

def add_field(par, instr, placeholder=''):
    """插入复杂域（TOC、PAGE 等）。"""
    def fld(kind):
        r = par.add_run()
        el = OxmlElement('w:fldChar')
        el.set(qn('w:fldCharType'), kind)
        r._r.append(el)
        return r
    fld('begin')
    r = par.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = instr
    r._r.append(it)
    fld('separate')
    if placeholder:
        run_fmt(par.add_run(placeholder), size=BODY_PT)
    fld('end')


def add_paragraph(doc, text):
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = Pt(BODY_PT * 2)
    add_inline(par, text)
    return par


def add_list_item(doc, marker, text, indent_level=0):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    left = 24 + indent_level * 18
    pf.left_indent = Pt(left)
    pf.first_line_indent = Pt(-18)
    add_inline(par, marker + ' ' + text)
    return par


def add_code_block(doc, lines):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.left_indent = Pt(6)
    pf.right_indent = Pt(6)
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F2F2F2')
    ppr.append(shd)
    bdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '4')
        el.set(qn('w:color'), 'BFBFBF')
        bdr.append(el)
    ppr.append(bdr)
    for k, line in enumerate(lines):
        run = par.add_run(line)
        run_fmt(run, size=CODE_PT, mono=True)
        if k < len(lines) - 1:
            run.add_break()
    return par


def add_image(doc, alt, path, base_dir):
    p = Path(path)
    if not p.is_absolute():
        p = base_dir / p
    if not p.exists():
        raise FileNotFoundError('图片不存在：%s' % p)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = bool(alt)
    par.add_run().add_picture(str(p), width=Cm(min(TEXT_WIDTH_CM, 14.5)))
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        cap.paragraph_format.space_after = Pt(8)
        add_inline(cap, alt, size=10.5)


# ---------------------------------------------------------------- 封面、目录、页码

def add_cover(doc, meta):
    def centered(text, size, bold=False, before=0, after=0):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = par.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        run_fmt(par.add_run(text), size=size, bold=bold)
        return par

    centered(meta.get('title', ''), 22, True, before=120, after=36)
    centered(meta.get('subtitle', ''), 18, True, after=150)
    for label, key in (('版本', 'version'), ('日期', 'date'), ('编制单位', 'org')):
        if meta.get(key):
            centered('%s：%s' % (label, meta[key]), 14, after=8)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(12)
    run_fmt(par.add_run('目　录'), size=16, bold=True)
    body = doc.add_paragraph()
    add_field(body, 'TOC \\o "1-3" \\h \\z \\u', '请在此处单击右键，选择“更新域”，生成目录。')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def setup_page(doc):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.8)
    sec.different_first_page_header_footer = True   # 封面不显示页码
    footer = sec.footer
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(par, 'PAGE', '1')
    for r in par.runs:
        run_fmt(r, size=9)


# ---------------------------------------------------------------- 主流程

def parse_front_matter(lines):
    meta = {}
    if lines and lines[0].strip() == '---':
        for k in range(1, len(lines)):
            if lines[k].strip() == '---':
                for ln in lines[1:k]:
                    if ':' in ln:
                        key, val = ln.split(':', 1)
                        meta[key.strip()] = val.strip()
                return meta, lines[k + 1:]
    return meta, lines


def convert(md_path, out_path, overrides=None):
    md_path = Path(md_path)
    lines = md_path.read_text(encoding='utf-8').splitlines()
    meta, lines = parse_front_matter(lines)
    meta.update({k: v for k, v in (overrides or {}).items() if v})

    doc = Document()
    setup_styles(doc)
    setup_page(doc)
    doc.core_properties.title = meta.get('subtitle') or meta.get('title', '')
    doc.core_properties.author = meta.get('org', '')
    if meta.get('title') or meta.get('subtitle'):
        add_cover(doc, meta)
        add_toc(doc)

    para_buf = []

    def flush_para():
        if para_buf:
            add_paragraph(doc, ''.join(para_buf))
            para_buf.clear()

    i, n = 0, len(lines)
    base_dir = md_path.parent
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush_para()
            i += 1
            continue
        if stripped.startswith('```'):
            flush_para()
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith('```'):
                code.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, code)
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            flush_para()
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
            i += 1
            continue
        m = re.match(r'^!\[(.*?)\]\((.+?)\)$', stripped)
        if m:
            flush_para()
            add_image(doc, m.group(1), m.group(2), base_dir)
            i += 1
            continue
        if stripped.startswith('|') and i + 1 < n and is_sep_row(lines[i + 1]):
            flush_para()
            rows = [split_row(lines[i])]
            i += 2
            while i < n and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        m = re.match(r'^(\s*)([-*+])\s+(.*)$', line)
        if m:
            flush_para()
            add_list_item(doc, '•', m.group(3), len(m.group(1)) // 2)
            i += 1
            continue
        m = re.match(r'^(\s*)(\d+)[.)]\s+(.*)$', line)
        if m:
            flush_para()
            add_list_item(doc, m.group(2) + '.', m.group(3), len(m.group(1)) // 2)
            i += 1
            continue
        para_buf.append(stripped)
        i += 1
    flush_para()
    doc.save(str(out_path))
    return out_path


def main():
    ap = argparse.ArgumentParser(description='Markdown 转宋体 Word 文档')
    ap.add_argument('input')
    ap.add_argument('output')
    ap.add_argument('--title')
    ap.add_argument('--subtitle')
    ap.add_argument('--version')
    ap.add_argument('--date')
    ap.add_argument('--org')
    a = ap.parse_args()
    convert(a.input, a.output, {'title': a.title, 'subtitle': a.subtitle, 'version': a.version,
                                'date': a.date, 'org': a.org})
    print('已生成', a.output)


if __name__ == '__main__':
    sys.exit(main())
