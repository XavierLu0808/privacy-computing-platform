#!/usr/bin/env python3
"""把 lo_update_toc.py 产出的目录条目写进 docx 的 TOC 域结果（域本身保留，Word 里仍可右键更新）。

用法：python3 tools/inject_toc.py 输入.docx toc.json 输出.docx [--shift N] [--calibrate 注入后渲染的.pdf]
  --shift N       所有页码加 N（目录自身页数与 LibreOffice 生成的目录不同时用来校正）
  --calibrate PDF 根据「已注入目录的文档」渲染出的 PDF，自动算出 --shift（各标题实际页码与目录页码之差的中位数）
"""
import re
import json
import sys

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

sys.path.insert(0, __file__.rsplit('/', 1)[0])
from md2docx import run_fmt, TEXT_WIDTH_CM  # noqa: E402


def fld_char(par, kind):
    r = par.add_run()
    el = OxmlElement('w:fldChar')
    el.set(qn('w:fldCharType'), kind)
    r._r.append(el)


def find_toc_paragraph(doc):
    for p in doc.paragraphs:
        if 'TOC \\o' in p._p.xml:
            return p
    raise RuntimeError('文档中没有 TOC 域')


def style_entry(par, level):
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_after = Pt(3 if level == 1 else 1)
    pf.space_before = Pt(6 if level == 1 else 0)
    pf.left_indent = Pt(12 * (level - 1))
    pf.first_line_indent = Pt(0)
    ppr = par._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:leader'), 'dot')
    tab.set(qn('w:pos'), str(int(TEXT_WIDTH_CM / 2.54 * 1440)))
    tabs.append(tab)
    ppr.append(tabs)


def calibrate(entries, pdf_path):
    """各标题在 PDF 中最后一次单独成行的页码（目录页在前，正文在后）与目录页码之差的中位数。"""
    import fitz
    texts = [p.get_text() for p in fitz.open(pdf_path)]
    diffs = []
    for _, title, page in entries:
        found = [i + 1 for i, x in enumerate(texts) if re.search(r'^' + re.escape(title) + r'\s*$', x, re.M)]
        if found:
            diffs.append(found[-1] - int(page))
    diffs.sort()
    return diffs[len(diffs) // 2] if diffs else 0


def inject(src, toc_json, dst, shift=0, pdf=None):
    entries = json.load(open(toc_json, encoding='utf-8'))
    if pdf:
        shift = calibrate(entries, pdf)
    entries = [[lv, t, str(int(pg) + shift)] for lv, t, pg in entries]
    doc = Document(src)
    first = find_toc_paragraph(doc)
    # 清掉占位结果：保留 begin / instr / separate 三个 run，删掉其余
    runs = first.runs
    keep = 0
    for i, r in enumerate(runs):
        if r._r.find(qn('w:fldChar')) is not None and r._r.find(qn('w:fldChar')).get(qn('w:fldCharType')) == 'separate':
            keep = i + 1
            break
    for r in runs[keep:]:
        r._r.getparent().remove(r._r)
    prev = first
    last = None
    for k, (level, title, page) in enumerate(entries):
        if k == 0:
            par = first
        else:
            new_p = OxmlElement('w:p')
            prev._p.addnext(new_p)
            from docx.text.paragraph import Paragraph
            par = Paragraph(new_p, first._parent)
        style_entry(par, level)
        size = 12 if level == 1 else 10.5
        run_fmt(par.add_run(title + '\t' + page), size=size, bold=(level == 1))
        prev = par
        last = par
    fld_char(last, 'end')
    doc.save(dst)
    return len(entries), shift


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument('src'); ap.add_argument('toc'); ap.add_argument('dst')
    ap.add_argument('--shift', type=int, default=0)
    ap.add_argument('--calibrate')
    a = ap.parse_args()
    n, s = inject(a.src, a.toc, a.dst, a.shift, a.calibrate)
    print('已写入目录条目 %d 条，页码校正 %+d' % (n, s))
