#!/usr/bin/env python3
"""md2docx 的单元测试：python3 -m unittest tools/test_md2docx.py"""
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
import md2docx  # noqa: E402

SAMPLE = """---
title: 示例项目
subtitle: 示例文档
version: V1.0
date: 2026 年 10 月
org: 复旦大学
---

# 1 引言

## 1.1 目的

这是一段正文，含 **加粗** 与 `code`。

1. 第一项
2. 第二项

- 无序项

| 名称 | 说明 |
|---|---|
| a\\|b | 含转义竖线 |
| x | y |

```
line1
line2
```
"""


class Md2DocxTest(unittest.TestCase):
    def build(self, text=SAMPLE):
        d = Path(tempfile.mkdtemp())
        src, out = d / 'a.md', d / 'a.docx'
        src.write_text(text, encoding='utf-8')
        md2docx.convert(src, out)
        return Document(str(out))

    def test_headings_and_cover(self):
        doc = self.build()
        heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith('Heading')]
        self.assertEqual(heads, [('Heading 1', '1 引言'), ('Heading 2', '1.1 目的')])
        texts = [p.text for p in doc.paragraphs]
        self.assertIn('示例项目', texts)
        self.assertIn('版本：V1.0', texts)
        self.assertIn('编制单位：复旦大学', texts)

    def test_fonts_are_songti(self):
        doc = self.build()
        rfonts = doc.styles['Normal'].element.rPr.find(qn('w:rFonts'))
        self.assertEqual(rfonts.get(qn('w:eastAsia')), '宋体')
        self.assertEqual(rfonts.get(qn('w:ascii')), 'Times New Roman')
        h1 = doc.styles['Heading 1'].element.rPr.find(qn('w:rFonts'))
        self.assertEqual(h1.get(qn('w:eastAsia')), '宋体')
        self.assertIsNone(h1.get(qn('w:asciiTheme')))
        for p in doc.paragraphs:
            for r in p.runs:
                f = r._element.rPr.find(qn('w:rFonts')) if r._element.rPr is not None else None
                if f is not None:
                    self.assertEqual(f.get(qn('w:eastAsia')), '宋体')

    def test_table(self):
        doc = self.build()
        self.assertEqual(len(doc.tables), 1)
        t = doc.tables[0]
        self.assertEqual([c.text for c in t.rows[0].cells], ['名称', '说明'])
        self.assertEqual(t.rows[1].cells[0].text, 'a|b')
        self.assertIn('w:tblBorders', t._element.tblPr.xml)
        self.assertIn('D9D9D9', t.rows[0].cells[0]._element.xml)
        self.assertIn('w:tblHeader', t.rows[0]._tr.xml)

    def test_lists_code_toc_and_page_field(self):
        doc = self.build()
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(any(t.startswith('1. 第一项') for t in texts))
        self.assertTrue(any(t.startswith('• 无序项') for t in texts))
        self.assertIn('line1\nline2', texts)
        body_xml = doc.element.body.xml
        self.assertIn('TOC \\o "1-3"', body_xml)
        footer_xml = doc.sections[0].footer._element.xml
        self.assertIn('PAGE', footer_xml)

    def test_image_embedded(self):
        from PIL import Image
        d = Path(tempfile.mkdtemp())
        Image.new('RGB', (40, 20), 'white').save(d / 'x.png')
        (d / 'a.md').write_text('# 1 图\n\n![图 1 示意](x.png)\n', encoding='utf-8')
        md2docx.convert(d / 'a.md', d / 'a.docx')
        doc = Document(str(d / 'a.docx'))
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertIn('图 1 示意', [p.text for p in doc.paragraphs])

    def test_missing_image_raises(self):
        with self.assertRaises(FileNotFoundError):
            self.build('# 1 图\n\n![图 1](nope.png)\n')


if __name__ == '__main__':
    unittest.main()
